import docx

doc = docx.Document('NutriVision_AI_PRD_1.docx')
print('Sections:', len(doc.sections))
for s in doc.sections:
    print('Page size:', s.page_width.pt, 'x', s.page_height.pt)
    print('Margins:', s.top_margin.pt, s.bottom_margin.pt, s.left_margin.pt, s.right_margin.pt)

print('Styles present:')
for st in doc.styles:
    if st.name in ['Normal', 'Heading 1', 'Heading 2', 'Heading 3', 'Table Grid', 'List Paragraph', 'Title']:
        f_name = st.font.name if hasattr(st, 'font') and st.font else None
        print(f'Style: {st.name}, font: {f_name}')

print('Tables count:', len(doc.tables))
for i, t in enumerate(doc.tables):
    print(f'Table {i}: {len(t.rows)} rows, {len(t.columns)} cols, style: {t.style.name if t.style else None}')
